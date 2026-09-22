"""数据视图：每个结果都与 pandas 独立重算的数字逐项对照；含故障注入与只读保证。"""

import json
import time
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd
import pyarrow.parquet as pq
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from dsflow.core.hashing import tree_manifest
from dsflow.core.project import Project
from dsflow.core.schemas import InvariantCheck
from dsflow.data import browse, compare, invariants
from dsflow.data.cache import DataCache
from dsflow.data.convert import retype_text_columns
from dsflow.data.datasets import DatasetStore
from dsflow.data.engine import DataError, QueryError, connect, lit
from dsflow.data.files import discover
from dsflow.data.profile import profile
from dsflow.data.xlsx_stream import iter_rows, list_sheets
from dsflow.index.db import PlatformIndex
from dsflow.jobs import JobManager
from dsflow.server.app import create_app

from .conftest import STEP_11

N = 5000


def frames():
    rng = np.random.default_rng(7)
    cat = rng.choice(["A", "B", "C"], N, p=[0.5, 0.3, 0.2]).astype(object)
    cat[rng.choice(N, 50, replace=False)] = None
    amount = rng.normal(100, 40, N).round(2)
    amount[rng.choice(N, 80, replace=False)] = np.nan
    a = pd.DataFrame({
        "id": np.arange(N),
        "cat": cat,
        "amount": amount,
        "qty": rng.integers(-2, 20, N),
        "ts": pd.Timestamp("2025-01-01") + pd.to_timedelta(rng.integers(0, 365 * 24, N), unit="h"),
        "code": [f"{i:05d}" for i in rng.integers(0, 3000, N)],
        "name": rng.choice(np.array(["螺栓", "螺母 ", "", "垫片", None], dtype=object), N),
    })
    drop = (a["cat"] == "C") & (a["id"] % 3 == 0)
    b = a[~drop].copy()
    changed = b.index[b["amount"].notna()][:7]
    b.loc[changed, "amount"] += 1
    b = b.drop(columns=["name"])
    b["flag"] = b["qty"] > 5
    return a, b, int(drop.sum()), sorted(b.loc[changed, "id"].tolist())


def write_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(["编码", "数量", "金额", "日期", "备注"])
    for i in range(300):
        amount = "未知" if i == 250 else i * 1.5
        ws.append([f"{i:05d}", i % 9, amount, pd.Timestamp("2025-01-01") + pd.Timedelta(days=i), None if i % 4 else "加急"])
    wb.save(path)


def write_inline_xlsx(path: Path) -> None:
    """模拟导出工具的写法：文字内嵌在单元格（inlineStr）、没有共享字符串表、单元格稀疏。"""
    ns = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
    rel = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
    rows = [
        '<row r="1"><c r="A1" t="inlineStr"><is><t>商品编号</t></is></c><c r="B1" t="inlineStr"><is><t>品牌</t></is></c>'
        '<c r="C1" t="inlineStr"><is><t>价格</t></is></c></row>',
        '<row r="2"><c r="A2" t="inlineStr"><is><t>00017</t></is></c><c r="C2"><v>12.5</v></c></row>',
        '<row r="3"><c r="A3" t="inlineStr"><is><t>00018</t></is></c><c r="B3" t="inlineStr"><is><t>正泰</t></is></c>'
        '<c r="C3"><v>3</v></c></row>',
    ]
    parts = {
        "[Content_Types].xml": '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>',
        "_rels/.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                       '<Relationship Id="rId1" Target="xl/workbook.xml" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument"/></Relationships>',
        "xl/workbook.xml": f'<workbook {ns} {rel}><sheets><sheet name="Sheet1" r:id="rId3" sheetId="1"/></sheets></workbook>',
        "xl/_rels/workbook.xml.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                                      '<Relationship Id="rId3" Target="worksheets/sheet1.xml" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"/></Relationships>',
        "xl/worksheets/sheet1.xml": f'<worksheet {ns}><sheetData>{"".join(rows)}</sheetData></worksheet>',
    }
    with zipfile.ZipFile(path, "w") as z:
        for name, text in parts.items():
            z.writestr(name, text)


@pytest.fixture
def data_project(tmp_path):
    root = tmp_path / "dp"
    (root / "data").mkdir(parents=True)
    (root / "lifecycle").mkdir()
    (root / "lifecycle" / "steps.json").write_text(json.dumps({
        "title": "数据测试", "stages": [{"id": 1, "name": "预处理", "dir": "steps/01"}], "steps": [],
        "dependencies": [],
    }, ensure_ascii=False), encoding="utf-8")
    a, b, dropped, changed = frames()
    a.to_parquet(root / "data" / "a.parquet", index=False)
    b.to_parquet(root / "data" / "b.parquet", index=False)
    a.to_csv(root / "data" / "a.csv", index=False)
    pd.DataFrame({"订单号": ["SO1", "SO2", "SO3"], "金额": [1.5, 2.0, None], "类目": ["电力物资", "MRO", "办公"]}) \
        .to_csv(root / "data" / "gbk.csv", index=False, encoding="gbk")
    write_xlsx(root / "data" / "goods.xlsx")
    write_inline_xlsx(root / "data" / "inline.xlsx")
    (root / "data" / "~$goods.xlsx").write_bytes(b"lock")
    return {"project": Project(root, readonly=False), "a": a, "b": b, "dropped": dropped, "changed": changed}


def prepared(ctx, rel):
    return DataCache(ctx["project"]).prepare(rel)


# ---------- 发现、指纹、转换 ----------

def test_discover_skips_lock_files(data_project):
    paths = {f["path"]: f["format"] for f in discover(data_project["project"].root)}
    assert paths == {"data/a.parquet": "parquet", "data/b.parquet": "parquet", "data/a.csv": "csv",
                     "data/gbk.csv": "csv", "data/goods.xlsx": "xlsx", "data/inline.xlsx": "xlsx"}


def test_fingerprint_is_cached_and_tracks_changes(data_project):
    cache = DataCache(data_project["project"])
    first = cache.fingerprint("data/a.csv")
    assert cache.known_fingerprint("data/a.csv") == first
    path = data_project["project"].root / "data" / "a.csv"
    path.write_text(path.read_text(encoding="utf-8") + "\n", encoding="utf-8")
    assert cache.known_fingerprint("data/a.csv") is None
    assert cache.fingerprint("data/a.csv") != first


def test_csv_and_gbk_csv_convert(data_project):
    parquet = prepared(data_project, "data/a.csv")
    assert browse.row_count(parquet) == N
    gbk = pq.read_table(prepared(data_project, "data/gbk.csv")).to_pandas()
    assert list(gbk.columns) == ["订单号", "金额", "类目"] and gbk["类目"].tolist() == ["电力物资", "MRO", "办公"]
    assert prepared(data_project, "data/a.csv") == parquet  # 第二次复用缓存


def test_xlsx_convert_keeps_codes_as_text(data_project):
    cols = {c["name"]: c for c in browse.describe(prepared(data_project, "data/goods.xlsx"))}
    assert cols["编码"]["kind"] == "text" and cols["数量"]["kind"] == "numeric" and cols["日期"]["kind"] == "temporal"
    table = pq.read_table(prepared(data_project, "data/goods.xlsx")).to_pandas()
    assert len(table) == 300 and table["编码"].iloc[17] == "00017"


def test_inline_string_xlsx_streams(data_project):
    path = data_project["project"].root / "data" / "inline.xlsx"
    assert list_sheets(path) == [("Sheet1", "xl/worksheets/sheet1.xml")]
    assert list(iter_rows(path)) == [["商品编号", "品牌", "价格"], ["00017", None, "12.5"], ["00018", "正泰", "3"]]
    table = pq.read_table(prepared(data_project, "data/inline.xlsx")).to_pandas()
    assert table["商品编号"].tolist() == ["00017", "00018"] and table["价格"].tolist() == [12.5, 3.0]
    assert table["品牌"].isna().tolist() == [True, False]


@pytest.mark.parametrize("name", ["goods.xlsx", "inline.xlsx"])
def test_fast_parser_matches_tree_parser(data_project, name):
    path = data_project["project"].root / "data" / name
    assert list(iter_rows(path, fast=True)) == list(iter_rows(path, fast=False))


def test_openpyxl_dates_become_timestamps_via_stream(data_project):
    """同一列混有文字时 DuckDB 直接读会失败，走流式解析；日期样式单元格必须还原成日期而不是 Excel 序号。"""
    rows = list(iter_rows(data_project["project"].root / "data" / "goods.xlsx"))
    assert rows[1][3] == "2025-01-01 00:00:00" and rows[251][2] == "未知"


def test_retype_rules(tmp_path):
    src = tmp_path / "t.parquet"
    pd.DataFrame({
        "codes": ["00123", "00456", None], "nums": ["1.5", "2", None],
        "when": ["2025-01-01 00:00:00", "2025-02-01 08:30:00", None], "mixed": ["1", "x", None],
    }).to_parquet(src, index=False)
    dst = tmp_path / "o.parquet"
    kept = retype_text_columns(connect(), src, dst)
    assert sorted(kept) == ["codes", "mixed"]
    kinds = {c["name"]: c["kind"] for c in browse.describe(dst)}
    assert kinds == {"codes": "text", "nums": "numeric", "when": "temporal", "mixed": "text"}


# ---------- 浏览、抽样、SQL ----------

def test_page_sort_and_filters_match_pandas(data_project):
    a = data_project["a"]
    parquet = prepared(data_project, "data/a.parquet")
    first = browse.page(parquet, offset=100, limit=50)
    assert first["total"] == N and [r[0] for r in first["rows"]] == list(range(100, 150))
    assert [r[1] for r in first["rows"]] == a["id"].iloc[100:150].tolist()

    top = browse.page(parquet, limit=5, sort="amount", desc=True)
    assert [r[3] for r in top["rows"]] == a["amount"].nlargest(5).tolist()
    last = browse.page(parquet, offset=N - 3, limit=3, sort="amount", desc=True)
    assert all(r[3] is None for r in last["rows"])  # 缺失值排在最后

    only_a = browse.page(parquet, filters=[{"column": "cat", "op": "eq", "value": "A"}])
    assert only_a["total"] == int((a["cat"] == "A").sum())
    assert browse.page(parquet, filters=[{"column": "amount", "op": "is_null"}])["total"] == int(a["amount"].isna().sum())
    contains = browse.page(parquet, filters=[{"column": "name", "op": "contains", "value": "螺"}])
    assert contains["total"] == int(a["name"].fillna("").str.contains("螺").sum())
    with pytest.raises(QueryError):
        browse.page(parquet, filters=[{"column": "nope", "op": "eq", "value": 1}])


def test_sample_is_reproducible(data_project):
    parquet = prepared(data_project, "data/a.parquet")
    s1, s2 = browse.sample(parquet, 30, seed=1), browse.sample(parquet, 30, seed=1)
    assert len(s1["rows"]) == 30 and s1["rows"] == s2["rows"]
    assert browse.sample(parquet, 30, seed=2)["rows"] != s1["rows"]


def test_sql_is_read_only_sandbox(data_project, tmp_path):
    parquet = prepared(data_project, "data/a.parquet")
    out = browse.run_sql(parquet, "SELECT cat, count(*) AS n FROM data GROUP BY cat ORDER BY cat NULLS LAST")
    counts = data_project["a"]["cat"].value_counts(dropna=False)
    assert [r[2] for r in out["rows"]] == [int(counts["A"]), int(counts["B"]), int(counts["C"]), int(counts[None])]
    secret = tmp_path / "secret.csv"
    secret.write_text("x\n1\n")
    for bad in (f"SELECT * FROM read_csv({lit(secret)})", "COPY data TO 'x.csv'", "SELECT 1; SELECT 2",
                "INSTALL httpfs", "SET enable_external_access=true"):
        with pytest.raises(QueryError):
            browse.run_sql(parquet, bad)
    assert browse.run_sql(parquet, "SELECT * FROM data", limit=10)["truncated"] is True
    with pytest.raises(QueryError, match="中止"):
        browse.run_sql(parquet, "SELECT count(*) FROM range(10000000000) a, range(100) b", timeout=0.3)


# ---------- 画像 ----------

def test_profile_matches_pandas(data_project):
    a = data_project["a"]
    rep = profile(prepared(data_project, "data/a.parquet"))
    cols = {c["name"]: c for c in rep["columns"]}
    assert rep["rows"] == N and rep["column_count"] == 7
    for name in a.columns:
        assert cols[name]["null_count"] == int(a[name].isna().sum()), name
    amount = cols["amount"]
    assert amount["mean"] == pytest.approx(a["amount"].mean(), rel=1e-9)
    assert (amount["min"], amount["max"]) == (a["amount"].min(), a["amount"].max())
    h = amount["histogram"]
    assert sum(h["counts"]) + h["under"] + h["over"] == int(a["amount"].notna().sum())
    assert cols["qty"]["negative"] == int((a["qty"] < 0).sum()) and cols["qty"]["zero"] == int((a["qty"] == 0).sum())
    assert {t["value"]: t["count"] for t in cols["cat"]["top"]} == a["cat"].value_counts().to_dict()
    assert cols["name"]["blank"] == int((a["name"].fillna("x").str.strip() == "").sum())
    assert sum(p["count"] for p in cols["ts"]["periods"]) == N
    assert cols["code"]["kind"] == "text" and cols["ts"]["kind"] == "temporal"


def test_heavy_tail_histogram_is_clipped(tmp_path):
    """少数极大值（如个别上亿的订单金额）不能让所有行挤进第一格：按 1%–99% 分位分箱，极端值另计。"""
    path = tmp_path / "t.parquet"
    values = list(np.linspace(0, 100, 1000)) + [1e8] * 5
    pd.DataFrame({"金额": values}).to_parquet(path, index=False)
    h = profile(path)["columns"][0]["histogram"]
    assert h["edges"][-1] < 1000 and h["over"] >= 5
    assert sum(h["counts"]) + h["under"] + h["over"] == 1005
    assert max(h["counts"]) < 0.2 * 1005  # 中间各格分布开了，而不是全挤在一格


def test_old_cached_profile_is_recomputed(data_project):
    cache = DataCache(data_project["project"])
    path = cache.profile_path("data/a.parquet")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps({"rows": 1, "columns": []}), encoding="utf-8")  # 没有 version 的旧格式
    assert cache.cached_profile("data/a.parquet") is None


# ---------- 对比 ----------

def test_compare_profiles_schema_and_rows(data_project):
    pa = profile(prepared(data_project, "data/a.parquet"))
    pb = profile(prepared(data_project, "data/b.parquet"))
    r = compare.compare_profiles(pa, pb)
    assert r["schema"]["added"] == ["flag"] and r["schema"]["removed"] == ["name"]
    assert r["rows"] == {"a": N, "b": N - data_project["dropped"], "delta": -data_project["dropped"],
                         "pct": -data_project["dropped"] / N}
    flagged = {c["name"]: c["flags"] for c in r["columns"]}
    # 删掉的行里没有 cat 为空的，所以 cat 的缺失数不变、不应被标记
    assert "均值变化" in flagged["amount"] and flagged["cat"] == []


def test_segment_impact_finds_concentration(data_project):
    a, b = data_project["a"], data_project["b"]
    r = compare.segment_impact(prepared(data_project, "data/a.parquet"), prepared(data_project, "data/b.parquet"), "cat")
    by = {row["value"]: row for row in r["rows"]}
    for value in ("A", "B", "C"):
        assert by[value]["a"] == int((a["cat"] == value).sum()) and by[value]["b"] == int((b["cat"] == value).sum())
    assert by["（空）"]["a"] == int(a["cat"].isna().sum())
    assert by["C"]["share_of_change"] == 1.0 and r["concentrated"] == ["C"]


def test_key_diff_matches_pandas(data_project):
    r = compare.key_diff(prepared(data_project, "data/a.parquet"), prepared(data_project, "data/b.parquet"), ["id"])
    assert (r["only_a"], r["only_b"], r["matched"]) == (data_project["dropped"], 0, N - data_project["dropped"])
    assert r["changed_rows"] == 7 and r["changed_by_column"] == {"amount": 7}
    assert sorted(s["key"]["id"] for s in r["sample_changed"]) == data_project["changed"][:5]
    assert all(s["changes"][0]["b"] == pytest.approx(s["changes"][0]["a"] + 1) for s in r["sample_changed"])


def test_key_diff_with_duplicate_keys_skips_column_diff(data_project):
    r = compare.key_diff(prepared(data_project, "data/a.parquet"), prepared(data_project, "data/b.parquet"), ["cat"])
    # cat 有 50 行为空（单独计数、不算重复），其余 4,950 行只有 3 个不同值
    assert r["null_keys_a"] == 50 and r["duplicates_a"] == (N - 50) - 3 and "changed_skipped" in r


def test_null_keys_are_excluded_not_matched(tmp_path):
    """主键为空的行不能按"空 = 空"配对：否则会凭空多出匹配、放大行数。"""
    a, b = tmp_path / "a.parquet", tmp_path / "b.parquet"
    pd.DataFrame({"k": [1, 2, None, 4], "v": [10, 20, 30, 40]}).to_parquet(a, index=False)
    pd.DataFrame({"k": [1, 2, None, 5], "v": [10, 21, 30, 50]}).to_parquet(b, index=False)
    r = compare.key_diff(a, b, ["k"])
    assert (r["null_keys_a"], r["null_keys_b"], r["duplicates_a"]) == (1, 1, 0)
    assert (r["matched"], r["only_a"], r["only_b"], r["changed_rows"]) == (2, 1, 1, 1)
    kept = invariants.run_check(a, b, InvariantCheck(type="keys_preserved", key=["k"]))
    assert not kept["passed"] and "1 行主键为空" in kept["detail"]


def test_distribution_uses_shared_bins(data_project):
    a, b = data_project["a"], data_project["b"]
    pa_, pb_ = prepared(data_project, "data/a.parquet"), prepared(data_project, "data/b.parquet")
    num = compare.distribution(pa_, pb_, "amount")
    assert sum(num["a"]) + num["under"][0] + num["over"][0] == int(a["amount"].notna().sum())
    assert sum(num["b"]) + num["under"][1] + num["over"][1] == int(b["amount"].notna().sum())
    assert len(num["edges"]) == len(num["a"]) + 1
    cat = compare.distribution(pa_, pb_, "cat")
    assert cat["kind"] == "categorical" and sum(cat["a"]) == N


# ---------- 不变量（含故障注入） ----------

def test_invariants_pass_and_fail(data_project):
    pa_, pb_ = prepared(data_project, "data/a.parquet"), prepared(data_project, "data/b.parquet")
    same = invariants.run_checks(pa_, pa_, [
        InvariantCheck(type="row_count_equal"),
        InvariantCheck(type="sum_equal", column="qty"),
        InvariantCheck(type="columns_unchanged", key=["id"], columns=["amount", "qty"]),
        InvariantCheck(type="keys_preserved", key=["id"]),
        InvariantCheck(type="no_new_nulls", columns=["cat"]),
    ])
    assert all(r["passed"] for r in same), same
    changed = {r["type"]: r for r in invariants.run_checks(pa_, pb_, [
        InvariantCheck(type="row_count_equal"),
        InvariantCheck(type="columns_unchanged", key=["id"], columns=["amount", "qty"]),
        InvariantCheck(type="columns_unchanged", key=["id"], columns=["qty"], description="数量原值不变"),
        InvariantCheck(type="keys_preserved", key=["id"]),
        InvariantCheck(type="row_count_max_change", tolerance=0.5),
        InvariantCheck(type="sum_equal", column="missing_col"),
    ])}
    assert not changed["row_count_equal"]["passed"]
    assert not changed["keys_preserved"]["passed"] and f"{data_project['dropped']:,}" in changed["keys_preserved"]["detail"]
    assert changed["row_count_max_change"]["passed"]
    assert not changed["sum_equal"]["passed"] and "无法检查" in changed["sum_equal"]["detail"]
    results = invariants.run_checks(pa_, pb_, [
        InvariantCheck(type="columns_unchanged", key=["id"], columns=["amount"]),
        InvariantCheck(type="columns_unchanged", key=["id"], columns=["qty"]),
    ])
    assert not results[0]["passed"] and "7 行" in results[0]["detail"] and results[1]["passed"]


# ---------- 数据集登记与只读 ----------

def test_dataset_versions_and_chain(data_project):
    store = DatasetStore(data_project["project"])
    first = store.register("data/a.parquet", "orders", "raw", description="一行 = 一个订单行")
    assert first["created"] and first["version"]["rows"] == N
    assert not store.register("data/a.parquet", "orders", "raw")["created"]
    store.register("data/b.parquet", "orders_clean", "processed", parents=["orders"], produced_by="1.2",
                   description="一行 = 一个清洗后的订单行")
    chain = store.chain()
    assert [c["name"] for c in chain] == ["orders", "orders_clean"] and chain[1]["parents"] == ["orders"]
    with pytest.raises(DataError):
        store.register("data/a.parquet", "bad name!", "raw")
    with pytest.raises(DataError):
        store.register("data/a.parquet", "x", "nowhere")


def test_readonly_project_gets_no_files(data_project, isolated_home):
    root = data_project["project"].root
    before = tree_manifest(root)
    project = Project(root, readonly=True)
    cache = DataCache(project)
    for rel in ("data/a.csv", "data/goods.xlsx"):
        cache.prepare(rel)
    DatasetStore(project).register("data/a.csv", "orders", "raw", description="一行 = 一个订单行")
    assert tree_manifest(root) == before
    assert cache.dir.is_relative_to(isolated_home)


# ---------- API ----------

def wait(client, job, timeout=60):
    deadline = time.time() + timeout
    while job["status"] in ("queued", "running"):
        assert time.time() < deadline, job
        time.sleep(0.05)
        job = client.get(f"/api/jobs/{job['id']}").json()
    assert job["status"] == "succeeded", job
    return job["result"]


def test_api_data_flow(data_project):
    client = TestClient(create_app(PlatformIndex(), JobManager()))
    pid = client.post("/api/projects", json={"path": str(data_project["project"].root), "readonly": False}).json()["id"]
    base = f"/api/projects/{pid}/data"

    overview = client.get(base).json()
    assert {f["path"] for f in overview["files"] if f["ready"]} == {"data/a.parquet", "data/b.parquet"}
    assert client.get(f"{base}/table", params={"path": "data/a.csv"}).status_code == 409

    assert wait(client, client.post(f"{base}/prepare", json={"path": "data/a.csv"}).json())["rows"] == N
    page = client.get(f"{base}/table", params={"path": "data/a.csv", "limit": 10, "sort": "id", "desc": True,
                                               "filters": json.dumps([{"column": "qty", "op": "ge", "value": 0}])}).json()
    assert page["rows"][0][1] == N - 1 and page["total"] == int((data_project["a"]["qty"] >= 0).sum())

    rep = wait(client, client.post(f"{base}/profile", json={"path": "data/a.parquet"}).json())
    assert rep["rows"] == N and client.get(f"{base}/profile", params={"path": "data/a.parquet"}).status_code == 200

    result = wait(client, client.post(f"{base}/compare", json={
        "a": "data/a.parquet", "b": "data/b.parquet", "key": ["id"], "segment": "cat"}).json())
    assert result["key"]["changed_rows"] == 7 and result["segment"]["concentrated"] == ["C"]

    checks = client.post(f"{base}/checks", json={"a": "data/a.parquet", "b": "data/b.parquet",
                                                  "checks": [{"type": "row_count_equal"}]}).json()
    assert checks[0]["passed"] is False
    assert client.post(f"{base}/sql", json={"path": "data/a.parquet", "sql": "DROP VIEW data"}).status_code == 400
    assert client.post(f"{base}/datasets", json={"path": "data/a.parquet", "name": "orders"}).status_code == 400, "不写介绍登记不进来"
    assert client.post(f"{base}/datasets", json={"path": "data/a.parquet", "name": "orders",
                                                 "description": "一行 = 一个订单行"}).status_code == 201
    assert client.get(base).json()["chain"][0]["name"] == "orders"
    assert client.get(f"{base}/table", params={"path": "../x.parquet"}).status_code == 400


def test_files_are_attributed_to_the_stage_that_owns_them(mini_project):
    """数据文件归哪个阶段：放在步骤目录下的算那一步，登记时写明由某一步产出的也算，其余的留空给项目总览。"""
    from dsflow.core.project import Project
    from dsflow.data.files import attribute

    from .conftest import STEP_11

    root, _ = mini_project
    reg, _ = Project(root, readonly=False).load()
    files = [{"path": f"{STEP_11}/outputs/核对清单.csv", "size": 10},
             {"path": "data/interim/clean.parquet", "size": 20},
             {"path": "data/raw/采购单.xlsx", "size": 30},
             {"path": "steps/02_EDA/画像.csv", "size": 40}]
    datasets = [{"name": "orders_clean", "versions": [
        {"path": "data/interim/clean.parquet", "stage": "processed", "produced_by": "2.1"}]}]
    got = {f["path"]: f for f in attribute(files, reg, datasets)}

    assert got[f"{STEP_11}/outputs/核对清单.csv"]["step"] == "1.1" and got[f"{STEP_11}/outputs/核对清单.csv"]["stage"] == 1
    assert got["data/interim/clean.parquet"]["step"] == "2.1" and got["data/interim/clean.parquet"]["stage"] == 2, "登记时写明产出步骤的中间文件跟着那一步走"
    assert got["data/interim/clean.parquet"]["dataset"] == "orders_clean" and got["data/interim/clean.parquet"]["data_stage"] == "processed"
    assert got["steps/02_EDA/画像.csv"]["stage"] == 2 and got["steps/02_EDA/画像.csv"]["step"] is None, "落在阶段目录、不在任何步骤目录下的算这个阶段"
    assert got["data/raw/采购单.xlsx"]["stage"] is None, "原始数据不属于某一个阶段，留在项目总览"


def test_data_overview_attributes_every_file(data_project):
    """总览接口回的每个文件都带着归属字段，网页据此分流到总览与各阶段。"""
    client = TestClient(create_app(PlatformIndex(), JobManager()))
    pid = client.post("/api/projects", json={"path": str(data_project["project"].root), "readonly": False}).json()["id"]
    files = client.get(f"/api/projects/{pid}/data").json()["files"]
    assert files and all("stage" in f and "step" in f and "dataset" in f for f in files)
    assert all(f["stage"] is None for f in files), "这个测试项目的数据都在 data/ 下，不属于任何阶段"


def test_data_query_answers_counts_and_refuses_writes(data_project):
    """答疑要能自己数数：一条只读 SQL 直接给出答案；改数据的语句一律回绝。"""
    from dsflow.protocol import call

    root = str(data_project["project"].root)
    prepared(data_project, "data/a.csv")
    got = call("data_query", {"project": root, "path": "data/a.csv", "sql": "SELECT count(*) AS 行数 FROM data"})
    assert got["ok"] and got["data"]["rows"][0][1] == N

    bad = call("data_query", {"project": root, "path": "data/a.csv", "sql": "DELETE FROM data"})
    assert bad["ok"] is False and "SELECT" in bad["error"]["hint"]

    not_ready = call("data_query", {"project": root, "path": "data/gbk.csv", "sql": "SELECT 1 FROM data"})
    assert not_ready["ok"] is False and "data_peek" in not_ready["error"]["hint"], "没转换的表要告诉模型改用 data_peek"


def test_peek_gives_every_column_name(data_project):
    """只看 8 列写不出 SQL：前几行照旧，但全部列名要给全。"""
    from dsflow.data.peek import peek

    got = peek(data_project["project"], "data/a.parquet", ["id"], 2)
    assert got["columns"] == ["id"], "要看的列照旧只给这几列"
    assert "qty" in got["all_columns"] and got["total_columns"] == len(got["all_columns"]) > 1


# ---------- 以数据表为主线看一个步骤 ----------


def test_step_tables_says_plainly_when_a_step_changes_no_data(mini_project):
    """校验型步骤不产出数据。要直说「没有改动」，不能留白让人以为漏了。"""
    from dsflow.data.datasets import DatasetStore
    from dsflow.data.steptables import step_tables

    root, _ = mini_project
    project = Project(root, readonly=False)
    reg, _ = project.load()
    path = root / "data" / "orders.csv"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("SKU,数量\nA,5\nB,1\n", encoding="utf-8")
    store = DatasetStore(project)
    store.register("data/orders.csv", "orders_raw", "raw", description="一行 = 一个订单行")
    store.link("orders_raw", "1.1", "核对")

    got = step_tables(project, reg, "1.1")
    assert got["changes_data"] is False
    assert [(t["name"], t["role"]) for t in got["tables"] if t["role"] == "核对"] == [("orders_raw", "核对")]


def test_step_tables_reports_the_columns_a_step_added(mini_project):
    from dsflow.data.datasets import DatasetStore
    from dsflow.data.steptables import step_tables

    root, _ = mini_project
    project = Project(root, readonly=False)
    reg, _ = project.load()
    data = root / "data"
    data.mkdir(parents=True, exist_ok=True)
    (data / "before.csv").write_text("SKU,数量\nA,5\n", encoding="utf-8")
    (data / "after.csv").write_text("SKU,数量,商品实体ID\nA,5,E-1\n", encoding="utf-8")
    store = DatasetStore(project)
    store.register("data/before.csv", "清洗后", "processed", description="一行 = 一个清洗后的订单行")
    store.register("data/after.csv", "权威明细", "processed", parents=["清洗后"], produced_by="1.1",
                   description="一行 = 一个订单行，1.1 补齐退货标记后得到")
    from dsflow.data.cache import DataCache

    cache = DataCache(project)
    for rel in ("data/before.csv", "data/after.csv"):
        cache.prepare(rel)

    got = step_tables(project, reg, "1.1")
    assert got["changes_data"] is True
    out = next(t for t in got["tables"] if t["role"] == "产出")
    assert out["diff"]["known"] and out["diff"]["added"] == ["商品实体ID"]
    assert out["diff"]["removed"] == [] and out["diff"]["kept"] == 2
    assert [t["name"] for t in got["tables"] if t["role"] == "输入"] == ["清洗后"]


def test_column_diff_says_it_cannot_tell_when_a_side_is_not_converted():
    from dsflow.data.steptables import diff_columns

    assert diff_columns(None, ["a"]) == {"known": False, "added": [], "removed": [], "kept": 0}


def test_step_tables_links_the_tables_the_guide_shows(mini_project):
    """讲解里配了数据视图的表，数据板块要能找到它并说明「讲解第几段讲到它」；没登记过的也补一张卡。"""
    from dsflow.data.datasets import DatasetStore
    from dsflow.data.steptables import step_tables
    from dsflow.explain.guide import put_guide

    root, _ = mini_project
    project = Project(root, readonly=False)
    reg, _ = project.load()
    data = root / "data"
    data.mkdir(parents=True, exist_ok=True)
    (data / "orders.csv").write_text("SKU,数量\nA,5\n", encoding="utf-8")
    (data / "ledger.csv").write_text("处理,行数\n原始,2\n", encoding="utf-8")
    (root / STEP_11 / "nb_1.1.ipynb").write_text(
        '{"cells":[{"cell_type":"code","metadata":{},"execution_count":1,"outputs":[],"source":"x = 1\n"}],'
        '"metadata":{},"nbformat":4,"nbformat_minor":5}', encoding="utf-8")
    DatasetStore(project).register("data/orders.csv", "orders_raw", "raw", description="一行 = 一个订单行")
    DatasetStore(project).link("orders_raw", "1.1", "核对")
    guide = {
        "step": "1.1", "notebook": "../../nb_1.1.ipynb",
        "brief": {"background": "b", "question": "q", "answer": "a", "can_continue": True, "did": "d", "next": "n"},
        "parts": [{"question": "台账对不对？", "cells": [{
            "cell": 1, "title": "看台账",
            "data": {"title": "这张台账本身", "frames": [
                {"label": "台账", "file": "data/ledger.csv", "columns": ["处理", "行数"], "caption": "加减对得上"},
                {"label": "源表", "file": "./data/orders.csv", "columns": ["SKU"]},
            ]},
        }]}],
    }
    import yaml as _yaml

    put_guide(project, reg, "1.1", _yaml.safe_dump(guide, allow_unicode=True))
    got = step_tables(project, reg, "1.1")
    by_name = {t["name"]: t for t in got["tables"]}
    assert by_name["orders_raw"]["role"] == "核对"
    assert [r["part"] for r in by_name["orders_raw"]["guide_refs"]] == [1]
    extra = by_name["这张台账本身"]
    assert extra["role"] == "讲解引用" and extra["path"] == "data/ledger.csv" and extra["note"] == "加减对得上"
    assert extra["guide_refs"][0]["question"] == "台账对不对？"
